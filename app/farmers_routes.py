# -*- coding: utf-8 -*-
"""Farmer Routes Blueprint for Chara Survey Automation."""
import os
import io
import json
import datetime
import urllib.request
import urllib.parse
from pathlib import Path
from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file, Response
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.db import (
    get_all_farmers,
    get_farmer_by_id,
    get_dashboard_counts,
    get_distinct_source_files,
    get_distinct_upload_batches,
    get_latest_batch_id,
    get_latest_source_file,
    update_farmer,
    verify_farmer,
    delete_farmer,
    get_connection,
)

farmers_bp = Blueprint("farmers", __name__)
FORM_SUBMIT_URL = "https://docs.google.com/forms/d/e/1FAIpQLSfh73QStzk0cZ6gvqRHZWG9V4yGsBv3tx2Z6P3WroXS97XxHQ/formResponse"

@farmers_bp.route("/farmers", methods=["GET"])
def list_farmers():
    status = request.args.get("status", "all")
    q = request.args.get("q", "").strip()
    source_file = request.args.get("source_file", "").strip() or None
    batch = request.args.get("batch", "").strip() or None
    latest = request.args.get("latest", "").lower() in ("true", "1", "yes")

    if latest or batch == "latest":
        latest_batch_id = get_latest_batch_id()
        batch = latest_batch_id

    farmers = get_all_farmers(
        status_filter=status,
        search_query=q,
        source_file=source_file,
        upload_batch_id=batch
    )
    counts = get_dashboard_counts(source_file=source_file, upload_batch_id=batch)

    source_files = get_distinct_source_files()
    upload_batches = get_distinct_upload_batches()
    latest_batch_id = get_latest_batch_id()

    return render_template(
        "farmers_dashboard.html",
        farmers=farmers,
        counts=counts,
        current_status=status,
        search_query=q,
        current_source_file=source_file or "all",
        current_batch=batch or "all",
        source_files=source_files,
        upload_batches=upload_batches,
        latest_batch_id=latest_batch_id
    )

@farmers_bp.route("/farmers/<int:farmer_id>/edit", methods=["GET", "POST"])
def edit_farmer_route(farmer_id: int):
    farmer = get_farmer_by_id(farmer_id)
    if not farmer:
        flash(f"नोंद #{farmer_id} सापडली नाही.", "danger")
        return redirect(url_for("farmers.list_farmers"))

    if request.method == "POST":
        def to_float(val, default=0.0):
            try:
                return float(val) if val not in (None, "") else default
            except (ValueError, TypeError):
                return default

        def to_int(val, default=0):
            try:
                return int(float(val)) if val not in (None, "") else default
            except (ValueError, TypeError):
                return default

        submit_action = request.form.get("submit_action", "save")
        status_val = "verified" if submit_action == "save_and_verify" else request.form.get("status", farmer["status"])

        update_data = {
            "taluka": request.form.get("taluka", "").strip(),
            "veterinary_hospital": request.form.get("veterinary_hospital", "").strip(),
            "revenue_circle": request.form.get("revenue_circle", "").strip(),
            "village_name": request.form.get("village_name", "").strip(),
            "farmer_name": request.form.get("farmer_name", "").strip(),
            "mobile": request.form.get("mobile", "").strip(),
            "cow_small": to_int(request.form.get("cow_small", 0)),
            "cow_large": to_int(request.form.get("cow_large", 0)),
            "buffalo_small": to_int(request.form.get("buffalo_small", 0)),
            "buffalo_large": to_int(request.form.get("buffalo_large", 0)),
            "milch_animals": to_int(request.form.get("milch_animals", 0)),
            "dry_animals": to_int(request.form.get("dry_animals", 0)),
            "daily_milk_litres": to_float(request.form.get("daily_milk_litres", 0.0)),
            "goats": to_int(request.form.get("goats", 0)),
            "sheep_migratory": to_int(request.form.get("sheep_migratory", 0)),
            "sheep_confined": to_int(request.form.get("sheep_confined", 0)),
            "green_fodder_tonnes": to_float(request.form.get("green_fodder_tonnes", 0.0)),
            "silage_tonnes": to_float(request.form.get("silage_tonnes", 0.0)),
            "dry_fodder_tonnes": to_float(request.form.get("dry_fodder_tonnes", 0.0)),
            "rabi_area_hectares": to_float(request.form.get("rabi_area_hectares", 0.0)),
            "water_availability": request.form.get("water_availability", "").strip(),
            "irrigation_available": request.form.get("irrigation_available", "").strip(),
            "irrigation_types": request.form.get("irrigation_types", "").strip(),
            "status": status_val,
            "review_notes": request.form.get("review_notes", "").strip()
        }

        # Check required fields
        if not update_data["farmer_name"]:
            flash("शेतकऱ्याचे नाव आवश्यक आहे.", "danger")
            return render_template("edit_farmer.html", farmer={**farmer, **update_data})

        success = update_farmer(farmer_id, update_data)
        if success:
            if status_val == "verified":
                flash(f"शेतकरी #{farmer_id} ({update_data['farmer_name']}) ची माहिती यशस्वीरित्या अद्यतनित केली व पडताळणी पूर्ण (Verified) झाली!", "success")
            else:
                flash(f"शेतकरी #{farmer_id} ({update_data['farmer_name']}) ची माहिती यशस्वीरित्या सेव्ह झाली!", "success")
            return redirect(url_for("farmers.list_farmers"))
        else:
            flash("नोंद सेव्ह करताना त्रुटी आली.", "danger")

    return render_template("edit_farmer.html", farmer=farmer)

@farmers_bp.route("/farmers/<int:farmer_id>/verify", methods=["POST"])
def verify_farmer_route(farmer_id: int):
    farmer = get_farmer_by_id(farmer_id)
    if not farmer:
        flash(f"नोंद #{farmer_id} सापडली नाही.", "danger")
        return redirect(url_for("farmers.list_farmers"))

    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    note = f"Manually verified on {now_str}"
    verify_farmer(farmer_id, note=note)
    flash(f"नोंद #{farmer_id} ({farmer['farmer_name']}) ची पडताळणी पूर्ण (Verified) झाली!", "success")
    return redirect(url_for("farmers.list_farmers"))

@farmers_bp.route("/farmers/<int:farmer_id>/delete", methods=["POST"])
def delete_farmer_route(farmer_id: int):
    farmer = get_farmer_by_id(farmer_id)
    name = farmer["farmer_name"] if farmer else f"#{farmer_id}"
    delete_farmer(farmer_id)
    flash(f"शेतकरी नोंद #{farmer_id} ({name}) सुरक्षितपणे हटवली.", "info")
    return redirect(url_for("farmers.list_farmers"))

@farmers_bp.route("/farmers/bulk_verify", methods=["POST"])
def bulk_verify_route():
    farmer_ids = request.form.getlist("farmer_ids")
    if not farmer_ids:
        flash("कृपया पडताळणी करण्यासाठी किमान एक नोंद निवडा.", "warning")
        return redirect(url_for("farmers.list_farmers"))

    verified_count = 0
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for f_id_str in farmer_ids:
        try:
            f_id = int(f_id_str)
            verify_farmer(f_id, note=f"Bulk verified on {now_str}")
            verified_count += 1
        except Exception:
            continue

    flash(f"निवडलेल्या {verified_count} शेतकरी नोंदींची पडताळणी यशस्वीरित्या पूर्ण झाली!", "success")
    return redirect(url_for("farmers.list_farmers"))

@farmers_bp.route("/farmers/export_excel", methods=["GET"])
def export_excel():
    status_filter = request.args.get("status", "all")
    source_file = request.args.get("source_file", "").strip() or None
    batch = request.args.get("batch", "").strip() or None

    farmers = get_all_farmers(
        status_filter=status_filter,
        source_file=source_file,
        upload_batch_id=batch
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "चारा सर्वेक्षण शेतकरी डेटा"

    headers = [
        "नोंद क्रमांक (ID)",
        "तालुका",
        "पशुवैद्यकीय दवाखाना",
        "महसूल मंडळ",
        "गावाचे नाव",
        "शेतकऱ्याचे नाव",
        "मोबाईल क्रमांक",
        "लहान गायी",
        "मोठ्या गायी",
        "लहान म्हशी",
        "मोठ्या म्हशी",
        "दुभती जनावरे",
        "भाकड जनावरे",
        "दूध (लिटर/दिवस)",
        "शेळ्या",
        "मेंढ्या",
        "पाण्याची उपलब्धता",
        "हिरवा चारा (टन)",
        "मुरघास (टन)",
        "सुका चारा (टन)",
        "रब्बी क्षेत्र (हेक्टर)",
        "सिंचन सोय",
        "सिंचन प्रकार",
        "डुप्लिकेट फ्लॅग",
        "डुप्लिकेट कारण",
        "स्थिती (Status)",
        "पडताळणी टिप्पणी",
        "मूळ फाईल",
        "नोंद दिनांक"
    ]
    ws.append(headers)

    # Style header row
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    # Append data rows
    for r in farmers:
        ws.append([
            r["id"],
            r["taluka"],
            r["veterinary_hospital"],
            r["revenue_circle"],
            r["village_name"],
            r["farmer_name"],
            str(r["mobile"]) if r["mobile"] else "",
            r["cow_small"],
            r["cow_large"],
            r["buffalo_small"],
            r["buffalo_large"],
            r["milch_animals"],
            r["dry_animals"],
            r["daily_milk_litres"],
            r["goats"],
            r["sheep_migratory"],
            r["water_availability"],
            r["green_fodder_tonnes"],
            r["silage_tonnes"],
            r["dry_fodder_tonnes"],
            r["rabi_area_hectares"],
            r["irrigation_available"],
            r["irrigation_types"],
            "होय (Duplicate)" if r["is_duplicate"] else "नाही",
            r["duplicate_reason"],
            r["status"],
            r["review_notes"],
            r["source_file"],
            r["created_at"]
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # Save to in-memory buffer
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"chara_survey_export_{status_filter}_{timestamp}.xlsx"

    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@farmers_bp.route("/farmers/<int:farmer_id>/submit_form", methods=["POST"])
def submit_to_form(farmer_id: int):
    # Mandatory Safety Rule: External form submission disabled by default to protect sensitive survey data
    enable_external = os.environ.get("ENABLE_EXTERNAL_FORM_SUBMISSION", "false").lower() in ("true", "1", "yes")
    if not enable_external:
        flash(
            "सुरक्षा सूचना: बाह्य गुगल फॉर्म सबमिशन सध्या अक्षम (Disabled) आहे. सर्व संवेदनशील डेटा स्थानिक SQLite डेटाबेसमध्ये सुरक्षित आहे. आवश्यक असल्यास ENABLE_EXTERNAL_FORM_SUBMISSION=true सेट करा.",
            "warning"
        )
        return redirect(url_for("farmers.list_farmers"))

    farmer = get_farmer_by_id(farmer_id)
    if not farmer:
        flash("शेतकरी नोंद सापडली नाही.", "danger")
        return redirect(url_for("farmers.list_farmers"))

    if farmer["status"] != "verified":
        flash(f"शेतकरी #{farmer_id} सबमिट करता येणार नाही: फक्त पडताळणी पूर्ण (Verified) नोंदीच सबमिट केल्या जाऊ शकतात.", "warning")
        return redirect(url_for("farmers.list_farmers"))

    with get_connection() as conn:
        already = conn.execute("SELECT id, submitted_at FROM form_submissions WHERE farmer_id = ?", (farmer_id,)).fetchone()
        if already:
            flash(f"शेतकरी #{farmer_id} आधीच दिनांक {already['submitted_at']} रोजी सबमिट केला गेला आहे.", "info")
            return redirect(url_for("farmers.list_farmers"))

    payload = {
        "entry.504627010": farmer["taluka"] or "नगर",
        "entry.329100902": farmer["veterinary_hospital"] or "Ralegun",
        "entry.1449625859": farmer["revenue_circle"] or "Walki",
        "entry.667616365": farmer["village_name"] or "",
        "entry.744608837": farmer["farmer_name"] or "",
        "entry.1412319850": str(farmer["mobile"] or ""),
        "entry.169726096": str(farmer["cow_small"] or 0),
        "entry.1315182064": str(farmer["cow_large"] or 0),
        "entry.1628192218": str(farmer["buffalo_small"] or 0),
        "entry.1401846099": str(farmer["buffalo_large"] or 0),
        "entry.1732260657": str(farmer["milch_animals"] or 0),
        "entry.647819049": str(farmer["dry_animals"] or 0),
        "entry.1725598946": str(farmer["daily_milk_litres"] or 0.0),
        "entry.540129639": str(farmer["goats"] or 0),
        "entry.241067737": str(farmer["sheep_migratory"] or 0),
        "entry.696685367": str(farmer["sheep_confined"] or 0),
        "entry.1850744836": farmer["water_availability"] or "उपलब्ध",
        "entry.355015442": farmer["green_fodder_available"] or "हिरवा",
        "entry.1035898857": str(farmer["green_fodder_tonnes"] or 0.0),
        "entry.901068256": str(farmer["silage_tonnes"] or 0.0),
        "entry.1545782155": str(farmer["dry_fodder_tonnes"] or 0.0),
        "entry.1534432430": str(farmer["fodder_for_sale_tonnes"] or 0.0),
        "entry.1686749785": str(farmer["rabi_area_hectares"] or 0.0),
        "entry.1359339260": farmer["irrigation_available"] or "आहे",
        "entry.509380968": farmer["irrigation_types"] or "विहीर",
        "entry.636337089": farmer["irrigation_other"] or ""
    }

    try:
        data = urllib.parse.urlencode(payload).encode("utf-8")
        req = urllib.request.Request(
            FORM_SUBMIT_URL,
            data=data,
            headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        )
        with urllib.request.urlopen(req, timeout=15) as resp:
            code = resp.getcode()
            snippet = resp.read()[:200].decode("utf-8", errors="ignore")

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        with get_connection() as conn:
            conn.execute(
                "INSERT INTO form_submissions (farmer_id, farmer_name, mobile, submitted_at, http_status, response_snippet) VALUES (?, ?, ?, ?, ?, ?)",
                (farmer_id, farmer["farmer_name"], farmer["mobile"], now_str, code, snippet)
            )
            update_farmer(farmer_id, {"status": "submitted"})

        flash(f"शेतकरी #{farmer_id} ({farmer['farmer_name']}) Google Form मध्ये यशस्वीरित्या सबमिट झाला!", "success")
    except Exception as e:
        flash(f"Google Form सबमिशन त्रुटी: {str(e)}", "danger")

    return redirect(url_for("farmers.list_farmers"))

@farmers_bp.route("/farmers/reextract", methods=["GET", "POST"])
def reextract_batch():
    """Re-extract an entire batch of records using the high-accuracy AI Vision engine."""
    batch = request.args.get("batch", "").strip()
    if not batch:
        flash("कृपया बॅच आयडी द्या.", "danger")
        return redirect(url_for("farmers.list_farmers"))

    with get_connection() as conn:
        cur = conn.cursor()
        cur.execute("SELECT id, source_file FROM farmers WHERE upload_batch_id = ? ORDER BY id ASC", (batch,))
        farmer_rows = cur.fetchall()

    if not farmer_rows:
        flash(f"बॅच {batch} मध्ये कोणत्याही नोंदी सापडल्या नाहीत.", "warning")
        return redirect(url_for("farmers.list_farmers"))

    source_file_name = farmer_rows[0]["source_file"]
    upload_dir = Path(__file__).resolve().parent.parent / "data" / "uploads"
    file_path = upload_dir / source_file_name

    if not file_path.exists():
        matches = list(upload_dir.glob(f"*{source_file_name}"))
        if matches:
            file_path = matches[0]

    if not file_path.exists():
        flash(f"मूळ फाईल '{source_file_name}' सर्व्हरवर सापडली नाही.", "danger")
        return redirect(url_for("farmers.list_farmers", batch=batch))

    try:
        from pipeline.extractor import extract_document
        res = extract_document(file_path)
        records = res.get("records", [])

        if not records:
            flash("AI व्हिजनद्वारे नोंदी काढता आल्या नाहीत.", "warning")
            return redirect(url_for("farmers.list_farmers", batch=batch))

        with get_connection() as conn:
            cur = conn.cursor()
            for idx, r in enumerate(records):
                if idx < len(farmer_rows):
                    f_id = farmer_rows[idx]["id"]
                    cur.execute("""
                        UPDATE farmers
                        SET farmer_name = ?,
                            mobile = ?,
                            village_name = ?,
                            taluka = ?,
                            daily_milk_litres = ?,
                            cow_large = ?,
                            cow_small = ?,
                            buffalo_large = ?,
                            buffalo_small = ?,
                            milch_animals = ?,
                            dry_animals = ?,
                            goats = ?,
                            green_fodder_tonnes = ?,
                            dry_fodder_tonnes = ?,
                            irrigation_available = ?,
                            irrigation_types = ?,
                            raw_ocr_text = ?,
                            uncertain_fields = ?,
                            status = ?
                        WHERE id = ?
                    """, (
                        r.get("farmer_name", ""),
                        r.get("mobile", ""),
                        r.get("village_name", ""),
                        r.get("taluka", ""),
                        r.get("daily_milk_litres", 0.0),
                        r.get("cow_large", 0),
                        r.get("cow_small", 0),
                        r.get("buffalo_large", 0),
                        r.get("buffalo_small", 0),
                        r.get("milch_animals", 0),
                        r.get("dry_animals", 0),
                        r.get("goats", 0),
                        r.get("green_fodder_tonnes", 0.0),
                        r.get("dry_fodder_tonnes", 0.0),
                        r.get("irrigation_available", "आहे"),
                        r.get("irrigation_types", "विहीर"),
                        r.get("raw_ocr_text", ""),
                        r.get("uncertain_fields", ""),
                        r.get("status", "pending_review"),
                        f_id
                    ))
            conn.commit()

        flash(f"यशस्वी! बॅच {batch} मधील सर्व नोंदी AI व्हिजनद्वारे नव्याने वाचून अचूक अपडेट केल्या!", "success")
    except Exception as e:
        flash(f"AI एक्स्ट्रॅक्शन करताना त्रुटी: {str(e)}", "danger")

    return redirect(url_for("farmers.list_farmers", batch=batch))

